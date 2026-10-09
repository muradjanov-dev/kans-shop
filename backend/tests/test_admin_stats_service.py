import json
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models.enums import OrderStatus, OrderType, PaymentMethod, PaymentStatus
from app.db.models.order import Order
from app.db.models.order_item import OrderItem
from app.db.models.user import User
from app.services import stats_service


async def _add_order(
    session: AsyncSession,
    user: User,
    *,
    order_number: str,
    created_at: datetime,
    total: str,
    status: OrderStatus = OrderStatus.NEW,
    payment_status: PaymentStatus = PaymentStatus.PENDING,
) -> Order:
    order = Order(
        order_number=order_number,
        user_id=user.id,
        order_type=OrderType.DELIVERY,
        status=status,
        customer_name="Synthetic Buyer",
        customer_phone="+998901234567",
        subtotal=Decimal(total),
        total=Decimal(total),
        payment_method=PaymentMethod.CASH,
        payment_status=payment_status,
        created_at=created_at,
    )
    session.add(order)
    await session.flush()
    session.add(
        OrderItem(
            order_id=order.id,
            product_name_snapshot="Synthetic pen",
            product_sku_snapshot="SYNTHETIC-PEN",
            price=Decimal(total),
            quantity=1,
            total=Decimal(total),
        )
    )
    await session.flush()
    return order


def _translator(language: str):
    locale_path = Path(__file__).parents[1] / "app" / "locales" / f"{language}.json"
    locale = json.loads(locale_path.read_text())

    def translate(key: str, **values: object) -> str:
        raw = locale["admin"][key.removeprefix("admin.")]
        return raw.format(**values)

    return translate


async def test_tashkent_period_and_amount_labels(
    db_session: AsyncSession, admin, user: User
) -> None:
    period_bounds = getattr(stats_service, "period_bounds", None)
    get_admin_stats = getattr(stats_service, "get_admin_stats", None)
    assert callable(period_bounds), "stats_service must export period_bounds"
    assert callable(get_admin_stats), "stats_service must export get_admin_stats"

    now = datetime(2026, 10, 9, 7, 0, tzinfo=UTC)
    today_start, today_end = period_bounds("today", now=now)
    assert today_start.isoformat() == "2026-10-08T19:00:00+00:00"
    assert today_end == now

    week_start, week_end = period_bounds("week", now=now)
    month_start, month_end = period_bounds("month", now=now)
    assert week_end - week_start == timedelta(days=7)
    assert month_end - month_start == timedelta(days=30)

    await _add_order(
        db_session,
        user,
        order_number="STAT-TASHKENT-1",
        created_at=datetime(2026, 10, 8, 20, 0, tzinfo=UTC),
        total="5000",
        payment_status=PaymentStatus.PAID,
    )
    stats = await get_admin_stats(db_session, admin_id=admin.id, period="today", now=now)
    assert stats.order_value == Decimal("5000")
    assert stats.revenue == Decimal("5000")  # legacy API alias keeps its meaning
    assert stats.paid_amount == Decimal("5000")

    for language in ("uz", "ru"):
        label = _translator(language)("admin.stats_order_value_label").casefold()
        assert "bekor qilinmagan" in label if language == "uz" else "без отмен" in label
        assert "tushum" not in label if language == "uz" else "выручка" not in label


async def test_cancelled_orders_count_but_not_order_value(
    db_session: AsyncSession, admin, user: User
) -> None:
    get_admin_stats = getattr(stats_service, "get_admin_stats", None)
    assert callable(get_admin_stats), "stats_service must export get_admin_stats"
    now = datetime(2026, 10, 9, 7, 0, tzinfo=UTC)
    await _add_order(
        db_session,
        user,
        order_number="STAT-CANCELLED-1",
        created_at=datetime(2026, 10, 8, 20, 0, tzinfo=UTC),
        total="20000",
        status=OrderStatus.CANCELLED,
    )
    await _add_order(
        db_session,
        user,
        order_number="STAT-ACTIVE-1",
        created_at=datetime(2026, 10, 9, 1, 0, tzinfo=UTC),
        total="10000",
    )

    stats = await get_admin_stats(db_session, admin_id=admin.id, period="today", now=now)
    assert stats.orders_count == 2
    assert stats.order_value == Decimal("10000")
    assert stats.avg_check == Decimal("10000")


async def test_paid_amount_requires_paid_and_non_cancelled(
    db_session: AsyncSession, admin, user: User
) -> None:
    get_admin_stats = getattr(stats_service, "get_admin_stats", None)
    assert callable(get_admin_stats), "stats_service must export get_admin_stats"
    now = datetime(2026, 10, 9, 7, 0, tzinfo=UTC)
    fixtures = [
        ("STAT-PAID-1", "5000", OrderStatus.NEW, PaymentStatus.PAID),
        ("STAT-PAID-CANCELLED", "20000", OrderStatus.CANCELLED, PaymentStatus.PAID),
        ("STAT-PENDING-1", "3000", OrderStatus.NEW, PaymentStatus.PENDING),
    ]
    for order_number, total, status, payment_status in fixtures:
        await _add_order(
            db_session,
            user,
            order_number=order_number,
            created_at=datetime(2026, 10, 8, 20, 0, tzinfo=UTC),
            total=total,
            status=status,
            payment_status=payment_status,
        )

    stats = await get_admin_stats(db_session, admin_id=admin.id, period="today", now=now)
    assert stats.paid_amount == Decimal("5000")
    assert stats.order_value == Decimal("8000")


async def test_xlsx_uses_same_period_and_labels(
    db_session: AsyncSession, admin, user: User
) -> None:
    export_admin_stats_xlsx = getattr(stats_service, "export_admin_stats_xlsx", None)
    assert callable(export_admin_stats_xlsx), "stats_service must export XLSX reports"
    now = datetime(2026, 10, 9, 7, 0, tzinfo=UTC)
    await _add_order(
        db_session,
        user,
        order_number="STAT-XLSX-1",
        created_at=datetime(2026, 10, 8, 20, 0, tzinfo=UTC),
        total="5000",
        payment_status=PaymentStatus.PAID,
    )

    data = await export_admin_stats_xlsx(
        db_session, admin_id=admin.id, period="today", now=now
    )
    sheet = load_workbook(BytesIO(data), data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    labels = {row[0]: row[1] for row in rows if row and row[0] is not None}
    assert labels["Boshlanish (UTC)"] == "2026-10-08 19:00"
    assert labels["Bekor qilinmagan buyurtmalar summasi (so'm)"] == 5000
    assert labels["To'langan summa (so'm)"] == 5000
