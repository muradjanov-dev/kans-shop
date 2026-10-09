import secrets
from collections.abc import AsyncIterator
from contextlib import asynccontextmanager
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from importlib import import_module
from importlib.util import find_spec
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import AsyncEngine

from app.core.config import settings
from app.db.models.admin import Admin
from app.db.models.admin_audit_event import AdminAuditEvent
from app.db.models.admin_session import AdminSession
from app.db.models.enums import AdminRole, OrderStatus, OrderType, PaymentMethod, PaymentStatus
from app.db.models.notification_outbox import NotificationOutbox
from app.db.models.order import Order
from app.db.models.traffic_source import TrafficSource
from app.db.models.user import User
from app.services.admin_session_service import create_admin_session

from .api_helpers import ApiCase, make_api_case

traffic_source_service = (
    import_module("app.services.traffic_source_service")
    if find_spec("app.services.traffic_source_service") is not None
    else None
)

ADMIN_COOKIE = "__Host-kans-admin"
WEBAPP_ORIGIN = f"{settings.webapp_url.split('://', 1)[0]}://{settings.webapp_url.split('://', 1)[1].split('/', 1)[0]}"


@asynccontextmanager
async def _admin_session(case: ApiCase, role: AdminRole) -> AsyncIterator[tuple[Admin, str]]:
    telegram_id = 9_100_000_000_000_000 + secrets.randbelow(100_000_000)
    async with case.session_maker() as session:
        admin = Admin(telegram_id=telegram_id, full_name="Reports Test Admin", role=role)
        session.add(admin)
        await session.flush()
        principal, raw_cookie = await create_admin_session(
            session, admin_id=admin.id, now=datetime.now(UTC)
        )
        await session.commit()
        admin_id = admin.id
        csrf_token = principal.csrf_token
    case.client.cookies.set(ADMIN_COOKIE, raw_cookie)
    try:
        async with case.session_maker() as session:
            live_admin = await session.get(Admin, admin_id)
            assert live_admin is not None
            yield live_admin, csrf_token
    finally:
        case.client.cookies.delete(ADMIN_COOKIE)
        async with case.session_maker() as session:
            await session.execute(
                AdminSession.__table__.delete().where(AdminSession.admin_id == admin_id)
            )
            await session.execute(Admin.__table__.delete().where(Admin.id == admin_id))
            await session.commit()


async def test_stats_operator_read_only(test_engine: AsyncEngine) -> None:
    async with (
        make_api_case(test_engine, base_url="https://testserver") as case,
        _admin_session(case, AdminRole.OPERATOR) as (_, csrf_token),
    ):
        order_id: int | None = None
        async with case.session_maker() as session:
            order = Order(
                order_number=f"HTTP-{secrets.token_hex(5)}",
                user_id=case.user_id,
                order_type=OrderType.DELIVERY,
                status=OrderStatus.NEW,
                customer_name="Synthetic Reports Buyer",
                customer_phone="+998901234567",
                subtotal=Decimal("12345"),
                total=Decimal("12345"),
                payment_method=PaymentMethod.CASH,
                payment_status=PaymentStatus.PAID,
                created_at=datetime.now(UTC) - timedelta(minutes=1),
            )
            session.add(order)
            await session.flush()
            order_id = order.id
            await session.commit()

        try:
            overview = await case.client.get("/api/v1/admin/stats/overview?period=week")
            export = await case.client.get("/api/v1/admin/stats/export.xlsx?period=week")
            write = await case.client.post(
                "/api/v1/admin/sources",
                headers={"Origin": WEBAPP_ORIGIN, "X-CSRF-Token": csrf_token},
                json={"name": "Operator campaign", "code": "operator"},
            )

            assert overview.status_code == 200
            assert overview.headers["cache-control"] == "private, no-store"
            overview_data = overview.json()
            assert Decimal(str(overview_data["revenue"])) == Decimal("12345")
            assert Decimal(str(overview_data["order_value"])) == Decimal("12345")
            assert Decimal(str(overview_data["paid_amount"])) == Decimal("12345")
            assert export.status_code == 200
            assert export.headers["cache-control"] == "private, no-store"
            assert write.status_code == 403
            assert write.json()["error"]["code"] == "ADMIN_ROLE_REQUIRED"

            russian_export = await case.client.get(
                "/api/v1/admin/stats/export.xlsx?period=week",
                headers={"Accept-Language": "ru"},
            )
            assert russian_export.status_code == 200
            russian_sheet = load_workbook(
                BytesIO(russian_export.content), data_only=True
            ).active
            russian_labels = {
                row[0]: row[1]
                for row in russian_sheet.iter_rows(values_only=True)
                if row and row[0] is not None
            }
            assert "Сумма оплаченных заказов без отмен (сум)" in russian_labels
            assert russian_labels["Период"] == "Неделя"
        finally:
            if order_id is not None:
                async with case.session_maker() as session:
                    await session.execute(delete(Order).where(Order.id == order_id))
                    await session.commit()


async def test_source_outbox_is_id_only_recipient_scoped_and_transactional(
    db_session, admin: Admin
) -> None:
    from app.services.traffic_source_service import create_source, update_source

    recipient = Admin(
        telegram_id=9_200_000_000_000_001,
        full_name="Source Reports Recipient",
        role=AdminRole.MANAGER,
        notifications_enabled=True,
    )
    muted_recipient = Admin(
        telegram_id=9_200_000_000_000_002,
        full_name="Muted Source Reports Recipient",
        role=AdminRole.MANAGER,
        notifications_enabled=False,
    )
    inactive_recipient = Admin(
        telegram_id=9_200_000_000_000_003,
        full_name="Inactive Source Reports Recipient",
        role=AdminRole.MANAGER,
        is_active=False,
        notifications_enabled=True,
    )
    db_session.add_all([recipient, muted_recipient, inactive_recipient])
    await db_session.flush()

    transaction = await db_session.begin_nested()
    source = await create_source(
        db_session, admin_id=admin.id, name="Outbox Test", code="outbox-test"
    )
    source_id = source.id
    created_audit = await db_session.scalar(
        select(AdminAuditEvent).where(
            AdminAuditEvent.resource_type == "traffic_source",
            AdminAuditEvent.resource_id == str(source_id),
            AdminAuditEvent.action == "traffic_source.create",
        )
    )
    assert created_audit is not None
    created_event = await db_session.scalar(
        select(NotificationOutbox).where(
            NotificationOutbox.payload_id == str(created_audit.id)
        )
    )
    assert created_event is not None
    assert created_event.event_type == "traffic_source.created"
    assert created_event.aggregate_id == str(source_id)
    assert created_event.recipient_admin_id == recipient.id
    assert created_event.recipient_user_id is None
    assert created_event.dedupe_key == (
        f"traffic-source:{created_audit.id}:admin:{recipient.id}"
    )

    await update_source(
        db_session,
        admin_id=admin.id,
        source_id=source_id,
        name=None,
        active=False,
    )
    updated_audit = await db_session.scalar(
        select(AdminAuditEvent).where(
            AdminAuditEvent.resource_type == "traffic_source",
            AdminAuditEvent.resource_id == str(source_id),
            AdminAuditEvent.action == "traffic_source.update",
        )
    )
    assert updated_audit is not None
    updated_event = await db_session.scalar(
        select(NotificationOutbox).where(
            NotificationOutbox.payload_id == str(updated_audit.id)
        )
    )
    assert updated_event is not None
    assert updated_event.event_type == "traffic_source.updated"
    assert updated_event.aggregate_id == str(source_id)
    assert updated_event.recipient_admin_id == recipient.id
    assert updated_event.recipient_user_id is None
    source_outbox_count = await db_session.scalar(
        select(func.count())
        .select_from(NotificationOutbox)
        .where(NotificationOutbox.aggregate_id == str(source_id))
    )
    assert source_outbox_count == 2

    await transaction.rollback()
    assert (
        await db_session.scalar(
            select(func.count())
            .select_from(TrafficSource)
            .where(TrafficSource.id == source_id)
        )
        == 0
    )
    assert (
        await db_session.scalar(
            select(func.count())
            .select_from(AdminAuditEvent)
            .where(
                AdminAuditEvent.resource_type == "traffic_source",
                AdminAuditEvent.resource_id == str(source_id),
            )
        )
        == 0
    )
    assert (
        await db_session.scalar(
            select(func.count())
            .select_from(NotificationOutbox)
            .where(NotificationOutbox.aggregate_id == str(source_id))
        )
        == 0
    )


async def test_traffic_source_routes_publish_stats_and_audit_mutations(
    test_engine: AsyncEngine,
) -> None:
    async with (
        make_api_case(test_engine, base_url="https://testserver") as case,
        _admin_session(case, AdminRole.MANAGER) as (admin, csrf_token),
    ):
        headers = {"Origin": WEBAPP_ORIGIN, "X-CSRF-Token": csrf_token}
        code = f"api-{secrets.token_hex(5)}"
        created = await case.client.post(
            "/api/v1/admin/sources",
            headers=headers,
            json={"name": "  Instagram  ", "code": code.upper()},
        )
        assert created.status_code == 201
        source = created.json()
        assert source["name"] == "Instagram"
        assert source["code"] == code
        assert source["is_active"] is True
        assert "private, no-store" in created.headers["cache-control"]

        duplicate = await case.client.post(
            "/api/v1/admin/sources",
            headers=headers,
            json={"name": "Duplicate", "code": code},
        )
        assert duplicate.status_code == 409
        assert duplicate.json()["error"]["code"] == "TRAFFIC_SOURCE_CODE_EXISTS"

        listed = await case.client.get("/api/v1/admin/sources?page=1&limit=100")
        assert listed.status_code == 200
        assert any(item["id"] == source["id"] for item in listed.json()["items"])

        detail = await case.client.get(f"/api/v1/admin/sources/{source['id']}")
        assert detail.status_code == 200
        body = detail.json()
        assert body["clicks"] == 0
        assert body["first_touch_users"] == 0
        assert body["orders_count"] == 0
        assert Decimal(body["order_value"]) == Decimal("0")
        expected_username = settings.bot_username.removeprefix("@").strip()
        assert body["bot_link"] == (f"https://t.me/{expected_username}?start=src_{code}")

        invalid = await case.client.post(
            "/api/v1/admin/sources",
            headers=headers,
            json={"name": "Invalid", "code": "has space"},
        )
        assert invalid.status_code == 422
        assert invalid.json()["error"]["code"] == "VALIDATION_ERROR"

        patched = await case.client.patch(
            f"/api/v1/admin/sources/{source['id']}",
            headers=headers,
            json={"name": "Instagram October", "active": False},
        )
        assert patched.status_code == 200
        assert patched.json()["name"] == "Instagram October"
        assert patched.json()["is_active"] is False
        missing = await case.client.get("/api/v1/admin/sources/999999999")
        assert missing.status_code == 404
        assert missing.json()["error"]["code"] == "TRAFFIC_SOURCE_NOT_FOUND"

        async with case.session_maker() as session:
            events = list(
                (
                    await session.scalars(
                        select(AdminAuditEvent).where(
                            AdminAuditEvent.actor_admin_id == admin.id,
                            AdminAuditEvent.resource_type == "traffic_source",
                            AdminAuditEvent.resource_id == str(source["id"]),
                        )
                    )
                ).all()
            )
            assert {event.action for event in events} == {
                "traffic_source.create",
                "traffic_source.update",
            }
            assert any(
                event.after_json is not None and event.after_json.get("source_code") == code
                for event in events
            )
            audit_ids = [event.id for event in events]
            await session.execute(
                delete(NotificationOutbox).where(
                    NotificationOutbox.payload_id.in_(
                        [str(event_id) for event_id in audit_ids]
                    )
                )
            )
            await session.execute(
                delete(AdminAuditEvent).where(AdminAuditEvent.id.in_(audit_ids))
            )
            await session.execute(
                delete(TrafficSource).where(TrafficSource.id == source["id"])
            )
            await session.commit()


async def test_source_first_touch_and_used_source_deactivation(
    db_session, admin: Admin, user: User
) -> None:
    source = TrafficSource(code="first-touch", name="First Touch")
    other = TrafficSource(code="later-link", name="Later Link")
    db_session.add_all([source, other])
    await db_session.flush()
    user.traffic_source_id = source.id
    db_session.add(
        Order(
            order_number="SOURCE-PAID-1",
            user_id=user.id,
            order_type=OrderType.DELIVERY,
            status=OrderStatus.NEW,
            customer_name="Synthetic Buyer",
            customer_phone="+998901234567",
            subtotal=Decimal("12000"),
            total=Decimal("12000"),
            payment_method=PaymentMethod.CASH,
            payment_status=PaymentStatus.PAID,
        )
    )
    db_session.add(
        Order(
            order_number="SOURCE-CANCELLED-1",
            user_id=user.id,
            order_type=OrderType.DELIVERY,
            status=OrderStatus.CANCELLED,
            customer_name="Synthetic Buyer",
            customer_phone="+998901234567",
            subtotal=Decimal("7000"),
            total=Decimal("7000"),
            payment_method=PaymentMethod.CASH,
            payment_status=PaymentStatus.PAID,
        )
    )
    await db_session.flush()

    get_source_stats = getattr(traffic_source_service, "get_source_stats", None)
    update_source = getattr(traffic_source_service, "update_source", None)
    assert callable(get_source_stats), "traffic_source_service must export get_source_stats"
    assert callable(update_source), "traffic_source_service must export update_source"

    stats = await get_source_stats(db_session, admin_id=admin.id, source_id=source.id)
    assert stats.clicks == 0
    assert stats.first_touch_users == 1
    assert stats.orders_count == 2
    assert stats.order_value == Decimal("12000")

    await update_source(
        db_session, admin_id=admin.id, source_id=source.id, name=None, active=False
    )
    await db_session.refresh(user)
    await db_session.refresh(source)
    assert not source.is_active
    assert user.traffic_source_id == source.id
    still_attributed = await get_source_stats(
        db_session, admin_id=admin.id, source_id=source.id
    )
    assert still_attributed.first_touch_users == 1

    audit = await db_session.scalar(
        select(AdminAuditEvent).where(
            AdminAuditEvent.actor_admin_id == admin.id,
            AdminAuditEvent.resource_type == "traffic_source",
            AdminAuditEvent.resource_id == str(source.id),
        )
    )
    assert audit is not None
    assert audit.action == "traffic_source.update"
